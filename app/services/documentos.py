"""Ashley document generation. User-scoped outputs; no external document converters."""
import io
import re
from pathlib import Path
from xml.sax.saxutils import escape
from docx import Document
from docx.shared import Cm, Pt
from pypdf import PdfReader
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, KeepTogether
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

MAX_TEXT=80000

def extract_text(path: Path):
    ext=path.suffix.lower()
    if ext=='.pdf':
        reader=PdfReader(str(path))
        if len(reader.pages)>250: raise ValueError('PDF muito longo (máximo 250 páginas).')
        text='\n'.join(p.extract_text() or '' for p in reader.pages)
    elif ext=='.docx':
        doc=Document(str(path))
        text='\n'.join(p.text for p in doc.paragraphs)
        for table in doc.tables:
            text+='\n'+'\n'.join(' | '.join(c.text for c in row.cells) for row in table.rows)
    elif ext in ('.txt','.md'):
        text=path.read_text(encoding='utf-8-sig')
    else: raise ValueError('Formato não suportado. Envie PDF, DOCX, TXT ou MD.')
    if not text.strip(): raise ValueError('Arquivo sem texto extraível. PDF digitalizado exige OCR, ainda não disponível.')
    if len(text)>MAX_TEXT: raise ValueError('Texto acima do limite de 80 mil caracteres. Divida o documento.')
    return text

def paragraphs(text):
    # Keep plain text as text; no HTML from the model is interpreted.
    return [x.strip() for x in re.split(r'\n\s*\n',text.strip()) if x.strip()]

def create_pair(text: str, title: str, base: Path):
    if not text.strip(): raise ValueError('Não há conteúdo para gerar documento.')
    if len(text)>MAX_TEXT: raise ValueError('Documento excede 80 mil caracteres.')
    base.parent.mkdir(parents=True,exist_ok=True)
    doc=Document(); sec=doc.sections[0];sec.top_margin=Cm(2.5);sec.bottom_margin=Cm(2.5)
    sec.left_margin=Cm(3);sec.right_margin=Cm(2)
    normal=doc.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(11)
    normal.paragraph_format.space_after=Pt(8)
    doc.add_heading(title,level=0)
    for p in paragraphs(text):
        doc.add_paragraph(p)
    doc.save(str(base.with_suffix('.docx')))
    font='Helvetica'; bold='Helvetica-Bold'
    for fontname,filename in [('AshleyDejaVu','/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'),('AshleyDejaVuBold','/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf')]:
        if Path(filename).exists() and fontname not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont(fontname,filename))
    if 'AshleyDejaVu' in pdfmetrics.getRegisteredFontNames(): font='AshleyDejaVu'
    if 'AshleyDejaVuBold' in pdfmetrics.getRegisteredFontNames(): bold='AshleyDejaVuBold'
    styles=getSampleStyleSheet()
    styles.add(ParagraphStyle(name='AshleyTitle',fontName=bold,fontSize=16,leading=22,spaceAfter=18,alignment=TA_CENTER))
    styles.add(ParagraphStyle(name='AshleyBody',fontName=font,fontSize=10.5,leading=16,spaceAfter=10,alignment=TA_JUSTIFY,wordWrap='CJK'))
    pdf=SimpleDocTemplate(str(base.with_suffix('.pdf')),pagesize=A4,rightMargin=57,leftMargin=85,topMargin=71,bottomMargin=71)
    story=[Paragraph(escape(title),styles['AshleyTitle'])]
    for p in paragraphs(text):
        story.append(Paragraph(escape(p).replace('\n','<br/>'),styles['AshleyBody']))
    pdf.build(story)
    return base.with_suffix('.pdf'),base.with_suffix('.docx')
